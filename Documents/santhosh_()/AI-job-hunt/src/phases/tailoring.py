"""Phase 2 — Resume Tailoring: parse JD, compare keywords, rewrite resume."""

import json
from pathlib import Path
from typing import Optional

from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.messages import HumanMessage, SystemMessage

from src.config import AppConfig, get_gemini_api_key, resolve_path
from src.models import Job, JobStatus, JDAnalysis, ResumeAnalysis
from src.memory import update_job_status
from src.tools.file_ops import ensure_application_dir


# ── JD Analysis with LLM ─────────────────────────────────────────────────

async def analyze_job_description(jd_text: str, config: AppConfig) -> JDAnalysis:
    """Extract structured requirements from a job description."""
    llm = ChatGoogleGenerativeAI(
        model=config.agent_settings.model,
        google_api_key=get_gemini_api_key(),
        max_output_tokens=2048,
    )

    system = SystemMessage(content=(
        "You are an expert ATS and resume optimization specialist. "
        "Analyze the job description and extract structured data. "
        "Return ONLY valid JSON with these keys: "
        "hard_skills (list of top 10 required technical skills/keywords), "
        "soft_skills (list of top 5 soft skills), "
        "seniority_level (junior/mid/senior), "
        "key_responsibilities (top 5 bullet points), "
        "required_tools (specific frameworks, platforms, tools), "
        "certifications (any mentioned certs). "
        "No markdown fences, just JSON."
    ))

    human = HumanMessage(content=f"Job Description:\n\n{jd_text}")
    response = await llm.ainvoke([system, human])

    try:
        data = json.loads(response.content)
        return JDAnalysis(**data)
    except (json.JSONDecodeError, TypeError):
        return JDAnalysis()


async def compare_resume_to_jd(
    resume_text: str,
    jd_analysis: JDAnalysis,
    config: AppConfig,
) -> ResumeAnalysis:
    """Compare base resume keywords against JD requirements."""
    llm = ChatGoogleGenerativeAI(
        model=config.agent_settings.model,
        google_api_key=get_gemini_api_key(),
        max_output_tokens=2048,
    )

    skills_list = ", ".join(jd_analysis.hard_skills + jd_analysis.required_tools)

    system = SystemMessage(content=(
        "You are an ATS resume optimization expert. Compare the resume text "
        "against the required keywords and return ONLY valid JSON with: "
        "keywords_present (list of JD keywords found in resume), "
        "keywords_missing (list of JD keywords NOT in resume), "
        "weak_bullets (list of resume bullet points that could be improved), "
        "keyword_match_pct (float 0-100), "
        "ats_score_estimate (float 0-100). "
        "No markdown fences."
    ))

    human = HumanMessage(content=(
        f"Required keywords/skills:\n{skills_list}\n\n"
        f"Resume text:\n{resume_text}"
    ))

    response = await llm.ainvoke([system, human])

    try:
        data = json.loads(response.content)
        return ResumeAnalysis(**data)
    except (json.JSONDecodeError, TypeError):
        return ResumeAnalysis()


# ── Resume Rewriting ──────────────────────────────────────────────────────

async def rewrite_resume(
    resume_text: str,
    job: Job,
    jd_analysis: JDAnalysis,
    resume_analysis: ResumeAnalysis,
    config: AppConfig,
) -> str:
    """Produce a tailored resume text optimized for the specific JD."""
    llm = ChatGoogleGenerativeAI(
        model=config.agent_settings.model,
        google_api_key=get_gemini_api_key(),
        max_output_tokens=4096,
    )

    system = SystemMessage(content=f"""You are an expert resume writer specializing in ATS optimization.
Rewrite the resume following these strict rules:

SUMMARY SECTION:
- 3-4 sentences max
- Open with the exact job title: "{job.role_title}"
- Mention the company name once: "{job.company_name}"
- Include the top 3 matched hard skills
- Quantify experience in years

SKILLS SECTION:
- List all matched hard skills first
- Remove skills not mentioned in the JD to keep the list focused
- Group skills by category (Languages, Frameworks, Cloud, Tools)

EXPERIENCE BULLET POINTS:
- Start every bullet with a strong action verb (Built, Designed, Led, etc.)
- Follow the formula: [Action] + [What] + [Result/Impact]
- Quantify at least 60% of bullets (%, $, x faster, N users)
- Incorporate JD keywords naturally — never keyword-stuff
- Highlight the most relevant project at the top of each role

FORMATTING:
- Maximum 1 page for < 5 years experience; 2 pages for 5+ years
- No tables, columns, graphics, or headers/footers (ATS compatibility)
- Use plain text formatting only

Missing keywords to incorporate naturally: {', '.join(resume_analysis.keywords_missing)}
Seniority level: {jd_analysis.seniority_level}

Return ONLY the rewritten resume text. No explanations or preamble.""")

    human = HumanMessage(content=(
        f"Original resume:\n{resume_text}\n\n"
        f"Job description:\n{job.job_description}\n\n"
        f"Key responsibilities:\n" +
        "\n".join(f"- {r}" for r in jd_analysis.key_responsibilities)
    ))

    response = await llm.ainvoke([system, human])
    return response.content


async def write_cover_letter(
    job: Job,
    jd_analysis: JDAnalysis,
    config: AppConfig,
) -> str:
    """Generate a tailored cover letter for the job."""
    llm = ChatGoogleGenerativeAI(
        model=config.agent_settings.model,
        google_api_key=get_gemini_api_key(),
        max_output_tokens=2048,
    )

    system = SystemMessage(content=f"""You are an expert cover letter writer.
Write a cover letter following this structure:

Paragraph 1: Why this role at {job.company_name} specifically
Paragraph 2: Top 2-3 achievements most relevant to the JD
Paragraph 3: How you'll contribute in the first 90 days
Closing: Professional and direct; include notice period: {config.user_profile.notice_period}

CONSTRAINTS:
- 250-350 words maximum
- Address it to "Hiring Manager" unless a name is given
- No generic filler — be specific to this job
- Mention the role title: "{job.role_title}"
- Reference at least 2 skills from: {', '.join(jd_analysis.hard_skills[:5])}

Return ONLY the cover letter text. No explanations.""")

    human = HumanMessage(content=(
        f"User: {config.user_profile.name}\n"
        f"Email: {config.user_profile.email}\n"
        f"Phone: {config.user_profile.phone}\n"
        f"LinkedIn: {config.user_profile.linkedin_url}\n\n"
        f"Job: {job.role_title} at {job.company_name}\n"
        f"Location: {job.location}\n\n"
        f"Key responsibilities:\n" +
        "\n".join(f"- {r}" for r in jd_analysis.key_responsibilities)
    ))

    response = await llm.ainvoke([system, human])
    return response.content


# ── DOCX Generation ──────────────────────────────────────────────────────

def save_resume_as_docx(resume_text: str, output_path: Path) -> Path:
    """Save the tailored resume as a .docx file."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    lines = resume_text.strip().split("\n")
    for line in lines:
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
            continue

        # Detect section headers (all-caps or ending with colon)
        if stripped.isupper() or (stripped.endswith(":") and len(stripped) < 60):
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(13)
        elif stripped.startswith("- ") or stripped.startswith("• "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            doc.add_paragraph(stripped)

    doc.save(str(output_path))
    return output_path


def save_resume_as_pdf(resume_text: str, output_path: Path) -> Path:
    """Save the tailored resume as a .pdf file using fpdf2."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Helvetica", size=11)

    for line in resume_text.strip().split("\n"):
        stripped = line.strip()
        if not stripped:
            pdf.ln(5)
            continue

        if stripped.isupper() or (stripped.endswith(":") and len(stripped) < 60):
            pdf.set_font("Helvetica", "B", 13)
            pdf.cell(0, 8, stripped, new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", size=11)
        elif stripped.startswith("- ") or stripped.startswith("• "):
            pdf.cell(5)
            pdf.multi_cell(0, 6, f"  {stripped}")
        else:
            pdf.multi_cell(0, 6, stripped)

    pdf.output(str(output_path))
    return output_path


# ── Main Tailoring Flow ──────────────────────────────────────────────────

async def run_tailoring(
    jobs: list[Job],
    config: AppConfig,
) -> list[dict]:
    """Execute the full resume tailoring phase for a list of approved jobs.

    Returns a list of result dicts with before/after scores.
    """
    # Load base resume
    resume_path = resolve_path(config.base_resume_path)
    if not resume_path.exists():
        return [{"error": f"Base resume not found at {resume_path}"}]

    if resume_path.suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(str(resume_path))
            resume_text = "\n".join(page.get_text() for page in doc)
        except ImportError:
            resume_text = resume_path.read_text(encoding="utf-8", errors="ignore")
    else:
        resume_text = resume_path.read_text(encoding="utf-8")

    results = []

    for job in jobs:
        try:
            # Step 1: Analyze JD
            jd_analysis = await analyze_job_description(job.job_description, config)

            # Step 2: Compare resume to JD
            resume_analysis = await compare_resume_to_jd(
                resume_text, jd_analysis, config
            )

            before_kw = resume_analysis.keyword_match_pct
            before_ats = resume_analysis.ats_score_estimate

            # Step 3: Rewrite resume
            tailored_resume = await rewrite_resume(
                resume_text, job, jd_analysis, resume_analysis, config
            )

            # Step 4: Write cover letter
            cover_letter = await write_cover_letter(job, jd_analysis, config)

            # Step 5: Save files
            app_dir = ensure_application_dir(job.company_name, job.role_title)
            docx_path = save_resume_as_docx(tailored_resume, app_dir / "resume_v1.docx")
            pdf_path = save_resume_as_pdf(tailored_resume, app_dir / "resume_v1.pdf")
            cl_path = app_dir / "cover_letter.txt"
            cl_path.write_text(cover_letter, encoding="utf-8")

            # Step 6: Re-analyze tailored resume
            after_analysis = await compare_resume_to_jd(
                tailored_resume, jd_analysis, config
            )
            after_kw = after_analysis.keyword_match_pct
            after_ats = after_analysis.ats_score_estimate

            # Step 7: Update DB
            update_job_status(
                job.job_id,
                JobStatus.TAILORED,
                resume_version=str(pdf_path),
                cover_letter_path=str(cl_path),
                keyword_match_before=before_kw,
                keyword_match_after=after_kw,
                ats_score_before=before_ats,
                ats_score_after=after_ats,
            )

            results.append({
                "job_id": job.job_id,
                "company": job.company_name,
                "role": job.role_title,
                "resume_path": str(pdf_path),
                "cover_letter_path": str(cl_path),
                "keyword_before": before_kw,
                "keyword_after": after_kw,
                "ats_before": before_ats,
                "ats_after": after_ats,
                "missing_keywords_before": len(resume_analysis.keywords_missing),
                "missing_keywords_after": len(after_analysis.keywords_missing),
            })

        except Exception as e:
            results.append({
                "job_id": job.job_id,
                "error": str(e),
            })

    return results


def format_tailoring_table(result: dict) -> str:
    """Format a single tailoring result as a markdown comparison table."""
    if "error" in result:
        return f"Error for {result.get('job_id', '?')}: {result['error']}"

    return f"""### {result['role']} at {result['company']}

| Metric              | Before | After |
|---------------------|--------|-------|
| Keyword match (%)   | {result['keyword_before']:.0f}%   | {result['keyword_after']:.0f}%  |
| ATS score estimate  | {result['ats_before']:.0f}    | {result['ats_after']:.0f}   |
| Missing keywords    | {result['missing_keywords_before']}     | {result['missing_keywords_after']}    |

Resume: `{result['resume_path']}`
Cover letter: `{result['cover_letter_path']}`
"""
